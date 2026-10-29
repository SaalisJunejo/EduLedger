"""Timetable endpoints.

- POST /api/v1/hod/timetable/upload - HOD-only ingestion of dated class
  occurrences from a multipart CSV or XLSX file.
- GET  /api/v1/timetable/classes    - HOD/ADMIN read-only listing of every
  scheduled occurrence (the institution-wide counterpart of the faculty's
  own GET /api/v1/faculty/classes).
"""

import csv
import io
from pathlib import Path
from zipfile import ZipFile

from flask import current_app, jsonify, request
from openpyxl import load_workbook
from sqlalchemy.exc import SQLAlchemyError

from ..extensions import db
from ..models import ScheduledClass, User
from ..services.auth import require_auth
from ..services.timetable import COLUMNS, timetable_timezone, validate_row
from ..util import utcnow
from . import api_v1
from .errors import api_error

_MAX_BYTES = 10 * 1024 * 1024
_MAX_ROWS = 5000
_MAX_XLSX_EXPANDED_BYTES = 50 * 1024 * 1024
# Cap on the read-only listing below (newest occurrences first).
_MAX_LISTED_CLASSES = 500


def _read_table(data: bytes, suffix: str) -> tuple[list[str], list[tuple[int, list]]]:
    """Read before creating any rows; reject malformed/oversized input atomically."""
    workbook = None
    try:
        if suffix == ".csv":
            rows = csv.reader(io.StringIO(data.decode("utf-8-sig"), newline=""), strict=True)
        else:
            with ZipFile(io.BytesIO(data)) as archive:
                if sum(member.file_size for member in archive.infolist()) > _MAX_XLSX_EXPANDED_BYTES:
                    raise ValueError("expanded XLSX exceeds 50 MB")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            rows = workbook.active.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError("file is empty; column headers are required")
        headers = [str(value).strip() if value is not None else "" for value in header]
        if any(not column for column in headers) or len(set(headers)) != len(headers):
            raise ValueError("column headers must be nonempty and unique")
        missing = sorted(set(COLUMNS) - set(headers))
        if missing:
            raise ValueError(f"missing columns: {', '.join(missing)}")
        records = []
        for number, row in enumerate(rows, start=2):
            if number > _MAX_ROWS + 1:
                raise ValueError(f"file exceeds {_MAX_ROWS} data rows")
            records.append((number, list(row)))
        return headers, records
    finally:
        if workbook is not None:
            workbook.close()


@api_v1.post("/hod/timetable/upload")
@require_auth(roles=["HOD"])
def upload_timetable():
    """Create valid rows and report spreadsheet row numbers/reasons for bad rows.

    Field: file. Required columns: course_code, faculty_employee_id, room,
    day_of_week (Monday-Sunday or 0-6), start_time, end_time (HH:MM[:SS]).
    Faculty department restriction is intentionally not enforced pending policy.
    Every upload creates dated occurrences; it is not an upsert/weekly series.
    """
    uploaded = request.files.get("file")
    if uploaded is None or not uploaded.filename:
        return api_error("timetable", "missing_file", "multipart file field is required", 400)
    suffix = Path(uploaded.filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        return api_error("timetable", "unsupported_file_type", "file must be .csv or .xlsx", 400)
    data = uploaded.stream.read(_MAX_BYTES + 1)
    if len(data) > _MAX_BYTES:
        return api_error("timetable", "file_too_large", "file must be at most 10 MB", 413)
    try:
        zone = timetable_timezone()
    except ValueError as exc:
        return api_error("timetable", "invalid_configuration", str(exc), 500)
    try:
        headers, records = _read_table(data, suffix)
    except Exception as exc:
        return api_error("timetable", "invalid_file", f"cannot parse timetable: {exc}", 400)

    created = []
    rejected = []
    now = utcnow()
    try:
        for number, values in records:
            if len(values) != len(headers):
                rejected.append({"row": number, "reasons": ["row has a different number of cells than the header"]})
                continue
            scheduled_class, reasons = validate_row(dict(zip(headers, values)), zone, now)
            if reasons:
                rejected.append({"row": number, "reasons": reasons})
            else:
                db.session.add(scheduled_class)
                created.append(scheduled_class)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        current_app.logger.exception("Timetable upload database failure")
        return api_error("timetable", "database_error", "timetable could not be saved; no rows were created", 500)
    return jsonify({
        "rows_created": len(created), "rows_rejected": rejected,
        "classes": [scheduled_class.to_dict() for scheduled_class in created],
        "timezone": current_app.config["TIMETABLE_TIMEZONE"],
    }), 201 if created else 200


@api_v1.get("/timetable/classes")
@require_auth(roles=["HOD", "ADMIN"])
def list_timetable_classes():
    """Every scheduled occurrence, newest start time first (at most 500).

    Query: optional ``professor_id`` to limit the list to one professor's
    occurrences. Each entry carries the scheduled professor's name so the
    HOD/Admin surfaces can offer a faculty filter without a second call.
    Faculty department restriction is intentionally not enforced pending
    policy, the same as the upload above.
    """
    professor_id = None
    raw_professor_id = (request.args.get("professor_id") or "").strip()
    if raw_professor_id:
        if not (
            raw_professor_id.isascii()
            and raw_professor_id.isdigit()
            and len(raw_professor_id) <= 10
            and 0 < int(raw_professor_id) <= 2147483647
        ):
            return api_error(
                "request", "invalid_professor_id", "professor_id must be a positive user id", 400
            )
        professor_id = int(raw_professor_id)

    query = db.session.query(ScheduledClass, User).outerjoin(
        User, User.id == ScheduledClass.professor_id
    )
    if professor_id is not None:
        query = query.filter(ScheduledClass.professor_id == professor_id)
    rows = query.order_by(
        ScheduledClass.start_time.desc(), ScheduledClass.id.desc(),
    ).limit(_MAX_LISTED_CLASSES).all()

    fields = ("id", "courseCode", "room", "professorId", "dayOfWeek", "startTime", "endTime")
    entries = []
    for scheduled_class, professor in rows:
        serialized = scheduled_class.to_dict()
        entry = {field: serialized[field] for field in fields}
        entry["professorName"] = professor.full_name if professor is not None else None
        entries.append(entry)
    return jsonify({"classes": entries, "count": len(entries), "limit": _MAX_LISTED_CLASSES})
