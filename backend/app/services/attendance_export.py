"""Monthly attendance export (Phase 1.3, requirements doc section 3.5).

The report is read from the `attendance_check_ins` mirror (Phase 1.2a), never
from the chain: one row per on-chain-locked check-in, carrying the lock's
transaction hash so any row can be verified against AttendanceLedger.

A ScheduledClass collects check-ins from two kinds of AttendanceSession:
plain attendance sessions (class_id = the ScheduledClass id as a string,
app/api/sessions.py) and faculty workload sessions (class_id
``workload-<id>``, app/api/faculty.py). Both are joined here and every row
says which one it came from. Check-ins made before the mirror existed have
no row and are therefore not in the report.

The month is a calendar month in TIMETABLE_TIMEZONE; dates and times in the
file are local to that zone too.
"""

import csv
import io
import re
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import String, cast, literal, or_

from ..extensions import db
from ..models import AttendanceCheckIn, AttendanceSession, ScheduledClass, StudentProfile, User
from .live_attendance import WORKLOAD_CLASS_PREFIX

FORMATS = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
}

COLUMNS = (
    "Date (local)",
    "Course code",
    "Room",
    "Session ID",
    "Session type",
    "Roll number",
    "Student name",
    "Status",
    "Check-in time (local)",
    "On-chain tx hash",
)
SUMMARY_COLUMNS = ("Roll number", "Student name", "Lectures attended", "Courses")
SUMMARY_SHEET = "Summary"

SESSION_TYPE_ATTENDANCE = "attendance"
SESSION_TYPE_WORKLOAD = "workload"
STATUS_PRESENT = "PRESENT"

_MONTH = re.compile(r"(\d{4})-(\d{2})")
# Spreadsheet apps evaluate cells starting with these as formulas.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def month_range_utc(value, zone: ZoneInfo) -> tuple[datetime, datetime]:
    """'YYYY-MM' -> [start, end) of that local calendar month, as naive UTC.

    Raises ValueError for anything that is not a real month.
    """
    match = _MONTH.fullmatch(str(value or "").strip())
    if match is None:
        raise ValueError("month must be given as YYYY-MM")
    year, month = int(match[1]), int(match[2])
    if not 1 <= month <= 12 or not 1970 <= year <= 9998:
        raise ValueError(f"{value} is not a valid month (YYYY-MM, 1970-01 or later)")
    next_year, next_month = (year + 1, 1) if month == 12 else (year, month + 1)
    start = datetime(year, month, 1, tzinfo=zone)
    end = datetime(next_year, next_month, 1, tzinfo=zone)
    return (
        start.astimezone(timezone.utc).replace(tzinfo=None),
        end.astimezone(timezone.utc).replace(tzinfo=None),
    )


def filename(month: str, fmt: str) -> str:
    """The Content-Disposition filename, e.g. eduledger-attendance-2026-03.xlsx."""
    return f"eduledger-attendance-{month}.{fmt}"


def build_report(
    start_utc: datetime,
    end_utc: datetime,
    zone: ZoneInfo,
    *,
    class_id: int | None = None,
    professor_id: int | None = None,
) -> tuple[list[tuple], list[tuple]]:
    """Every mirrored check-in in [start_utc, end_utc) -> (rows, summary).

    `class_id` limits the report to one ScheduledClass, `professor_id` to the
    classes that professor is scheduled for (a class they only substituted
    for is not theirs). Rows follow COLUMNS, sorted by date, course and roll
    number; summary rows follow SUMMARY_COLUMNS, one per student sorted by
    roll number, counting the lectures they attended (one per check-in, i.e.
    per distinct session) and how many distinct course codes those span.
    """
    class_key = cast(ScheduledClass.id, String)
    query = (
        db.session.query(AttendanceCheckIn, AttendanceSession, ScheduledClass, User, StudentProfile)
        .join(AttendanceSession, AttendanceCheckIn.session_id == AttendanceSession.id)
        .join(
            ScheduledClass,
            or_(
                AttendanceSession.class_id == class_key,
                AttendanceSession.class_id == literal(WORKLOAD_CLASS_PREFIX) + class_key,
            ),
        )
        .outerjoin(User, User.id == AttendanceCheckIn.student_id)
        .outerjoin(StudentProfile, StudentProfile.user_id == AttendanceCheckIn.student_id)
        .filter(
            AttendanceCheckIn.checked_in_at >= start_utc,
            AttendanceCheckIn.checked_in_at < end_utc,
        )
    )
    if class_id is not None:
        query = query.filter(ScheduledClass.id == class_id)
    if professor_id is not None:
        query = query.filter(ScheduledClass.professor_id == professor_id)

    keyed_rows = []
    students = {}
    for check_in, session, scheduled_class, user, profile in query.all():
        local = check_in.checked_in_at.replace(tzinfo=timezone.utc).astimezone(zone)
        date_text = local.strftime("%Y-%m-%d")
        time_text = local.strftime("%H:%M:%S")
        roll_number = profile.roll_number if profile is not None and profile.roll_number else ""
        name = (
            user.full_name
            if user is not None and user.full_name
            else f"Student {check_in.student_id}"
        )
        session_type = (
            SESSION_TYPE_WORKLOAD
            if session.class_id.startswith(WORKLOAD_CLASS_PREFIX)
            else SESSION_TYPE_ATTENDANCE
        )
        row = (
            date_text,
            scheduled_class.course_code,
            scheduled_class.room,
            session.id,
            session_type,
            roll_number,
            name,
            STATUS_PRESENT,
            time_text,
            check_in.onchain_tx_hash or "",
        )
        sort_key = (date_text, scheduled_class.course_code, roll_number, time_text, session.id)
        keyed_rows.append((sort_key, row))

        entry = students.setdefault(
            check_in.student_id,
            {"roll_number": roll_number, "name": name, "sessions": set(), "courses": set()},
        )
        # One check-in per student per session (the mirror's unique key), so
        # the session set counts the lectures attended, not the classes.
        entry["sessions"].add(session.id)
        entry["courses"].add(scheduled_class.course_code)

    rows = [row for _, row in sorted(keyed_rows, key=lambda item: item[0])]
    summary = [
        (entry["roll_number"], entry["name"], len(entry["sessions"]), len(entry["courses"]))
        for entry in sorted(students.values(), key=lambda e: (e["roll_number"], e["name"]))
    ]
    return rows, summary


def _write_sheet(sheet, header: tuple, rows: list[tuple]) -> None:
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
        # openpyxl turns strings starting with "=" into formulas; names and
        # rooms are user input, so every text cell is pinned to plain text.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    for index, title in enumerate(header, start=1):
        width = max([len(str(title))] + [len(str(row[index - 1])) for row in rows])
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 70)


def build_xlsx(rows: list[tuple], summary: list[tuple]) -> bytes:
    """Sheet "Attendance" (the rows) + sheet "Summary" (one row per student).

    The Summary row counts the student's lectures attended (check-ins) and
    the number of distinct courses those lectures belong to.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    _write_sheet(sheet, COLUMNS, rows)
    _write_sheet(workbook.create_sheet(SUMMARY_SHEET), SUMMARY_COLUMNS, summary)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _csv_safe(value):
    """Neutralize spreadsheet formula injection in user-supplied text."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return f"'{value}"
    return value


def build_csv(rows: list[tuple]) -> bytes:
    """The attendance rows only, UTF-8 with a BOM so Excel reads names correctly."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    for row in rows:
        writer.writerow([_csv_safe(value) for value in row])
    return buffer.getvalue().encode("utf-8-sig")
