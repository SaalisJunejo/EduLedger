"""End-to-end test of the Phase 1.3 monthly attendance export.

Run it against a live stack (from backend/):

    .\\venv\\Scripts\\python.exe test_export.py

Prerequisites:
    - Flask backend running (backend/: python run.py), same backend/.env as
      this script (expected local times use its TIMETABLE_TIMEZONE)
    - No chain needed: the export reads only the attendance_check_ins mirror,
      so this script writes its own mirror rows (fake but well-formed tx
      hashes) instead of running real on-chain scans.

Steps:
     0. preflight          - GET /health
     1. auth setup         - register + log in unique faculty A, faculty B, an
                             HOD (seeded into the whitelist) and two students
     2. seed               - faculty A: class A1 (two plain attendance
                             sessions) and class A2, a second course, with two
                             workload sessions ("workload-<id>") - one for the
                             month's lecture, one for the 00:30 local check-in,
                             since the mirror allows one row per session and
                             student; faculty B: class B1. Check-ins in
                             TEST_MONTH, plus one at 00:30 local on the 1st of
                             the next month (still the previous month in UTC)
     3. xlsx (all classes) - faculty A, no class_id: parses with openpyxl,
                             exact headers, exact rows (both session types,
                             sorted, B's class and next month's row excluded)
     4. Summary sheet      - student 1 (two sessions of one class) counts 2
                             lectures / 1 course; student 2 (two classes, two
                             courses) counts 2 lectures / 2 courses
     5. csv                - the same rows as the xlsx, filename .csv
     6. one class          - class_id=A1 -> only A1's rows
     7. local month        - next month's export holds the 00:30 local check-in
     8. other faculty      - faculty B exports A1 -> 403 professor/not_class_professor
     9. student            - 403 auth/insufficient_role
    10. bad month          - 400 request/invalid_month
    11. bad format         - 400 request/invalid_format
    12. empty month        - 200, header rows only (xlsx both sheets, and csv)
    13. timetable list     - HOD GET /timetable/classes sees every class with
                             its professor's name, newest first
    14. professor filter   - ?professor_id= returns only that professor's rows
    15. list guards        - FACULTY and STUDENT -> 403, no token -> 401
    16. export scope       - HOD ?professor_id=A exports A's rows only,
                             ?professor_id=B only B's
    17. scope conflicts    - class_id of the other professor -> 400
                             class_professor_mismatch; a non-numeric or
                             unknown professor_id -> 400 invalid_professor_id
    18. faculty scope      - faculty A with B's id -> 403; with their own id
                             the export is unchanged
    Cleanup removes this run's users, profiles, classes, sessions, check-ins
    and the HOD whitelist row.
"""

import csv
import io
import os
import sys
from datetime import datetime, timedelta, timezone
from uuid import uuid4
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook

os.environ["SCHEDULER_ENABLED"] = "false"

from app import create_app
from app.extensions import db
from app.models import (
    AttendanceCheckIn,
    AttendanceSession,
    FacultyProfile,
    InstitutionWhitelist,
    ScheduledClass,
    StudentProfile,
    User,
    generate_nonce,
)
from app.util import utcnow

# Windows consoles often default to cp1252, which cannot print the marks below.
if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE_URL = "http://127.0.0.1:5000/api/v1"
RUN_ID = uuid4().hex[:10]
PASSWORD = "ExportTestPass123!"
FACULTY_A_EMAIL = f"faculty.export.a.{RUN_ID}@gcu.edu.pk"
FACULTY_B_EMAIL = f"faculty.export.b.{RUN_ID}@gcu.edu.pk"
HOD_EMAIL = f"hod.export.{RUN_ID}@gcu.edu.pk"
STUDENT_1_EMAIL = f"student.export.1.{RUN_ID}@gcu.edu.pk"
STUDENT_2_EMAIL = f"student.export.2.{RUN_ID}@gcu.edu.pk"
ALL_EMAILS = [FACULTY_A_EMAIL, FACULTY_B_EMAIL, HOD_EMAIL, STUDENT_1_EMAIL, STUDENT_2_EMAIL]
FACULTY_A_NAME = "Dr. Export A"
FACULTY_B_NAME = "Dr. Export B"
HOD_NAME = "Dr. Export HOD"
HOD_EMPLOYEE_ID = f"EMP-EXP-HOD-{RUN_ID}"
DEPARTMENT = "Computer Science"
STUDENT_1_NAME = "Export Student One"
STUDENT_2_NAME = "Export Student Two"
ROLL_1 = f"ROLL-EXP-{RUN_ID}-1"
ROLL_2 = f"ROLL-EXP-{RUN_ID}-2"
# Two course codes: the Summary sheet counts lectures and distinct courses.
COURSE_CODE = f"TEST-EXPORT-{RUN_ID}"
COURSE_CODE_2 = f"TEST-EXPORT-{RUN_ID}-B"
COURSE_CODES = [COURSE_CODE, COURSE_CODE_2]

TEST_MONTH = "2025-03"
NEXT_MONTH = "2025-04"
EMPTY_MONTH = "1999-01"

CLASS_FIELDS = {
    "id",
    "courseCode",
    "room",
    "professorId",
    "professorName",
    "dayOfWeek",
    "startTime",
    "endTime",
}

HEADERS = [
    "Date (local)",
    "Course code",
    "Room",
    "Session ID",
    "Session type",
    "Roll number",
    "Student name",
    "Status",
    "Check-in time (local)",
    "On-chain tx hash",
]
SUMMARY_HEADERS = ["Roll number", "Student name", "Lectures attended", "Courses"]

APP = create_app()
ZONE = ZoneInfo(APP.config["TIMETABLE_TIMEZONE"])
AUTH = {"faculty_a": None, "faculty_b": None, "hod": None, "student_1": None}
STATE = {
    "faculty_a_id": None,
    "faculty_b_id": None,
    "hod_id": None,
    "student_1_id": None,
    "student_2_id": None,
    "class_ids": [],
    "a1_id": None,
    "b1_id": None,
    "missing_user_id": None,  # a user id that does not exist
    "xlsx_rows": None,
    "summary": None,
    "expected": None,  # TEST_MONTH rows for faculty A, in export order
    "expected_a1": None,
    "expected_b1": None,
    "expected_next": None,
}


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def api(method: str, path: str, token: str | None = None, **kwargs) -> tuple[int, dict]:
    """Call the API; returns (status, decoded JSON body). Never raises on HTTP errors."""
    headers = kwargs.pop("headers", {})
    if token:
        headers["Authorization"] = f"Bearer {token}"
    response = requests.request(method, f"{BASE_URL}{path}", headers=headers, timeout=15, **kwargs)
    try:
        payload = response.json()
    except ValueError:
        payload = {"raw_body": response.text}
    return response.status_code, payload


def export(token: str, **params) -> requests.Response:
    """GET /attendance/export, returning the raw response (files are not JSON)."""
    return requests.get(
        f"{BASE_URL}/attendance/export",
        headers={"Authorization": f"Bearer {token}"},
        params=params,
        timeout=30,
    )


def error_of(response: requests.Response) -> dict:
    try:
        return response.json()
    except ValueError:
        return {"raw_body": response.text[:200]}


def assert_file(response: requests.Response, month: str, ext: str) -> None:
    assert response.status_code == 200, (
        f"export returned {response.status_code}: {error_of(response)}"
    )
    disposition = response.headers.get("Content-Disposition", "")
    expected_name = f"eduledger-attendance-{month}.{ext}"
    assert "attachment" in disposition and expected_name in disposition, (
        f"Content-Disposition {disposition!r} does not name {expected_name}"
    )


def xlsx_sheets(response: requests.Response) -> dict[str, list[list[str]]]:
    """Every sheet as a list of rows of strings (None -> '')."""
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    sheets = {}
    for sheet in workbook.worksheets:
        sheets[sheet.title] = [
            ["" if value is None else str(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    workbook.close()
    return sheets


def csv_rows(response: requests.Response) -> list[list[str]]:
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def local_to_utc(text: str) -> datetime:
    """A local wall time in TIMETABLE_TIMEZONE -> naive UTC (the DB format)."""
    local = datetime.fromisoformat(text).replace(tzinfo=ZONE)
    return local.astimezone(timezone.utc).replace(tzinfo=None)


def fake_tx_hash() -> str:
    return f"0x{uuid4().hex}{uuid4().hex}"


def register_and_login(kind: str, email: str, body: dict) -> dict:
    status, payload = api("POST", f"/auth/register/{kind}", json={**body, "email": email, "password": PASSWORD})
    assert status == 201, f"{kind} registration ({email}) returned {status}: {payload}"
    status, payload = api("POST", "/auth/login", json={"email": email, "password": PASSWORD})
    assert status == 200, f"login {email} failed ({status}): {payload}"
    return payload


# --------------------------------------------------------------------------
# Steps
# --------------------------------------------------------------------------
def step_preflight():
    status, payload = api("GET", "/health")
    assert status == 200, f"/health returned {status}: {payload}"
    assert payload.get("status") == "ok", f"health status is not 'ok': {payload}"
    assert payload["dependencies"]["database_configured"], "database is not configured"
    print(f"   timezone={APP.config['TIMETABLE_TIMEZONE']} run_id={RUN_ID}")


def step_auth_setup():
    for key, email, employee_id, name in (
        ("faculty_a", FACULTY_A_EMAIL, f"EMP-EXP-A-{RUN_ID}", FACULTY_A_NAME),
        ("faculty_b", FACULTY_B_EMAIL, f"EMP-EXP-B-{RUN_ID}", FACULTY_B_NAME),
    ):
        login = register_and_login(
            "faculty",
            email,
            {"employee_id": employee_id, "full_name": name, "department": DEPARTMENT},
        )
        AUTH[key] = login["token"]
        STATE[f"{key}_id"] = login["user"]["id"]

    # HOD registration only accepts whitelisted staff, so the directory row
    # is written directly (the same fixture pattern as test_record_proposal).
    with APP.app_context():
        db.session.add(InstitutionWhitelist(
            employee_id=HOD_EMPLOYEE_ID, email=HOD_EMAIL,
            name=HOD_NAME, department=DEPARTMENT,
        ))
        db.session.commit()
    login = register_and_login(
        "hod",
        HOD_EMAIL,
        {"employee_id": HOD_EMPLOYEE_ID, "full_name": HOD_NAME, "department": DEPARTMENT},
    )
    AUTH["hod"] = login["token"]
    STATE["hod_id"] = login["user"]["id"]

    for key, email, name, roll in (
        ("student_1", STUDENT_1_EMAIL, STUDENT_1_NAME, ROLL_1),
        ("student_2", STUDENT_2_EMAIL, STUDENT_2_NAME, ROLL_2),
    ):
        login = register_and_login(
            "student",
            email,
            {"full_name": name, "roll_number": roll, "batch": "2022-2026", "department": DEPARTMENT},
        )
        AUTH[key] = login["token"]
        STATE[f"{key}_id"] = login["user"]["id"]
    print(
        f"   faculty A={STATE['faculty_a_id']} B={STATE['faculty_b_id']} "
        f"hod={STATE['hod_id']} students={STATE['student_1_id']},{STATE['student_2_id']}"
    )


def step_seed():
    """Classes, finished sessions and mirrored check-ins, written directly."""
    s1, s2 = STATE["student_1_id"], STATE["student_2_id"]
    with APP.app_context():
        now = utcnow()

        def scheduled(professor_id: int, room: str, start: str, course_code: str = COURSE_CODE) -> ScheduledClass:
            start_utc = local_to_utc(start)
            row = ScheduledClass(
                professor_id=professor_id, room=room, course_code=course_code,
                start_time=start_utc, end_time=start_utc + timedelta(hours=1),
                day_of_week=None,
            )
            db.session.add(row)
            return row

        def ended_session(class_key: str) -> AttendanceSession:
            row = AttendanceSession(
                class_id=class_key, current_nonce=generate_nonce(), nonce_expires_at=now,
                is_active=False, ended_at=now,
            )
            db.session.add(row)
            return row

        a1 = scheduled(STATE["faculty_a_id"], "EXPORT-A1", f"{TEST_MONTH}-10T09:00:00")
        a2 = scheduled(STATE["faculty_a_id"], "EXPORT-A2", f"{TEST_MONTH}-12T10:30:00", COURSE_CODE_2)
        b1 = scheduled(STATE["faculty_b_id"], "EXPORT-B1", f"{TEST_MONTH}-11T09:00:00")
        db.session.flush()
        STATE["class_ids"] = [a1.id, a2.id, b1.id]
        STATE["a1_id"] = a1.id
        STATE["b1_id"] = b1.id

        session_a1 = ended_session(str(a1.id))
        session_a1_second = ended_session(str(a1.id))  # second lecture, same class
        session_a2 = ended_session(f"workload-{a2.id}")
        # attendance_check_ins is unique per (session, student), so the
        # month-boundary check-in needs its own session: same class, same
        # session type, a later lecture.
        session_a2_next = ended_session(f"workload-{a2.id}")
        session_b1 = ended_session(str(b1.id))
        db.session.flush()

        def check_in(session: AttendanceSession, student_id: int, local: str) -> str:
            tx_hash = fake_tx_hash()
            db.session.add(AttendanceCheckIn(
                session_id=session.id, student_id=student_id,
                checked_in_at=local_to_utc(local), onchain_tx_hash=tx_hash,
            ))
            return tx_hash

        # Student 1: two sessions of the SAME class (2 lectures, 1 course).
        # Student 2: one session each of A1 and A2 (2 lectures, 2 courses).
        tx_1a = check_in(session_a1, s1, f"{TEST_MONTH}-10T09:15:00")
        tx_2a = check_in(session_a1, s2, f"{TEST_MONTH}-10T09:16:30")
        tx_2b = check_in(session_a2, s2, f"{TEST_MONTH}-12T11:00:00")
        tx_1c = check_in(session_a1_second, s1, f"{TEST_MONTH}-13T09:15:00")
        tx_b1 = check_in(session_b1, s2, f"{TEST_MONTH}-11T09:20:00")  # faculty B's class
        tx_next = check_in(session_a2_next, s2, f"{NEXT_MONTH}-01T00:30:00")  # local next month
        db.session.commit()
        STATE["missing_user_id"] = User.query.order_by(User.id.desc()).first().id + 1000

        row_1a = [f"{TEST_MONTH}-10", COURSE_CODE, "EXPORT-A1", str(session_a1.id), "attendance",
                  ROLL_1, STUDENT_1_NAME, "PRESENT", "09:15:00", tx_1a]
        row_2a = [f"{TEST_MONTH}-10", COURSE_CODE, "EXPORT-A1", str(session_a1.id), "attendance",
                  ROLL_2, STUDENT_2_NAME, "PRESENT", "09:16:30", tx_2a]
        row_2b = [f"{TEST_MONTH}-12", COURSE_CODE_2, "EXPORT-A2", str(session_a2.id), "workload",
                  ROLL_2, STUDENT_2_NAME, "PRESENT", "11:00:00", tx_2b]
        row_1c = [f"{TEST_MONTH}-13", COURSE_CODE, "EXPORT-A1", str(session_a1_second.id), "attendance",
                  ROLL_1, STUDENT_1_NAME, "PRESENT", "09:15:00", tx_1c]
        row_next = [f"{NEXT_MONTH}-01", COURSE_CODE_2, "EXPORT-A2", str(session_a2_next.id), "workload",
                    ROLL_2, STUDENT_2_NAME, "PRESENT", "00:30:00", tx_next]
        row_b1 = [f"{TEST_MONTH}-11", COURSE_CODE, "EXPORT-B1", str(session_b1.id), "attendance",
                  ROLL_2, STUDENT_2_NAME, "PRESENT", "09:20:00", tx_b1]
        STATE["expected"] = [row_1a, row_2a, row_2b, row_1c]
        STATE["expected_a1"] = [row_1a, row_2a, row_1c]
        STATE["expected_b1"] = [row_b1]
        STATE["expected_next"] = [row_next]
        print(
            f"   classes={STATE['class_ids']} sessions A1={session_a1.id},"
            f"{session_a1_second.id} A2(workload)={session_a2.id},"
            f"{session_a2_next.id} B1={session_b1.id}"
        )


def step_xlsx_all_classes():
    response = export(AUTH["faculty_a"], month=TEST_MONTH, format="xlsx")
    assert_file(response, TEST_MONTH, "xlsx")
    sheets = xlsx_sheets(response)
    assert list(sheets) == ["Attendance", "Summary"], f"unexpected sheets: {list(sheets)}"
    rows = sheets["Attendance"]
    assert rows[0] == HEADERS, f"unexpected headers: {rows[0]}"
    assert rows[1:] == STATE["expected"], (
        f"unexpected rows:\n   got      {rows[1:]}\n   expected {STATE['expected']}"
    )
    STATE["xlsx_rows"] = rows
    STATE["summary"] = sheets["Summary"]
    print(f"   {len(rows) - 1} rows; student 1 row: {rows[1]}")


def step_summary_sheet():
    summary = STATE["summary"]
    assert summary[0] == SUMMARY_HEADERS, f"unexpected Summary headers: {summary[0]}"
    # Student 1: two sessions of one class -> 2 lectures, 1 course.
    # Student 2: two classes with two course codes -> 2 lectures, 2 courses.
    expected = [
        [ROLL_1, STUDENT_1_NAME, "2", "1"],
        [ROLL_2, STUDENT_2_NAME, "2", "2"],
    ]
    assert summary[1:] == expected, f"unexpected Summary rows: {summary[1:]} (expected {expected})"
    print(f"   {summary[1]} / {summary[2]}")


def step_csv_same_rows():
    response = export(AUTH["faculty_a"], month=TEST_MONTH, format="csv")
    assert_file(response, TEST_MONTH, "csv")
    rows = csv_rows(response)
    assert rows == STATE["xlsx_rows"], (
        f"csv rows differ from the xlsx:\n   csv  {rows}\n   xlsx {STATE['xlsx_rows']}"
    )
    print(f"   {len(rows) - 1} rows, identical to the xlsx")


def step_single_class():
    response = export(AUTH["faculty_a"], month=TEST_MONTH, format="csv", class_id=STATE["a1_id"])
    assert_file(response, TEST_MONTH, "csv")
    rows = csv_rows(response)
    assert rows[0] == HEADERS, f"unexpected headers: {rows[0]}"
    assert rows[1:] == STATE["expected_a1"], f"unexpected class A1 rows: {rows[1:]}"
    print(f"   class {STATE['a1_id']}: {len(rows) - 1} rows")


def step_local_month_boundary():
    response = export(AUTH["faculty_a"], month=NEXT_MONTH, format="csv")
    assert_file(response, NEXT_MONTH, "csv")
    rows = csv_rows(response)
    assert rows[1:] == STATE["expected_next"], f"unexpected {NEXT_MONTH} rows: {rows[1:]}"
    print(f"   {NEXT_MONTH}-01 00:30 local is in {NEXT_MONTH}, not {TEST_MONTH}")


def step_other_faculty_forbidden():
    response = export(AUTH["faculty_b"], month=TEST_MONTH, class_id=STATE["a1_id"])
    payload = error_of(response)
    assert response.status_code == 403, f"expected 403, got {response.status_code}: {payload}"
    assert payload.get("check") == "professor", f"unexpected check: {payload}"
    assert payload.get("error") == "not_class_professor", f"unexpected error: {payload}"
    print(f"   rejected: {payload['message']}")


def step_student_forbidden():
    response = export(AUTH["student_1"], month=TEST_MONTH)
    payload = error_of(response)
    assert response.status_code == 403, f"expected 403, got {response.status_code}: {payload}"
    assert payload.get("check") == "auth", f"unexpected check: {payload}"
    assert payload.get("error") == "insufficient_role", f"unexpected error: {payload}"
    print(f"   rejected: {payload['message']}")


def step_bad_month():
    for month in ("2025-13", "March", "", "2025-3"):
        response = export(AUTH["faculty_a"], month=month)
        payload = error_of(response)
        assert response.status_code == 400, f"month={month!r}: expected 400, got {response.status_code}: {payload}"
        assert payload.get("check") == "request", f"month={month!r}: unexpected check: {payload}"
        assert payload.get("error") == "invalid_month", f"month={month!r}: unexpected error: {payload}"
    print(f"   rejected: {payload['message']}")


def step_bad_format():
    response = export(AUTH["faculty_a"], month=TEST_MONTH, format="pdf")
    payload = error_of(response)
    assert response.status_code == 400, f"expected 400, got {response.status_code}: {payload}"
    assert payload.get("check") == "request", f"unexpected check: {payload}"
    assert payload.get("error") == "invalid_format", f"unexpected error: {payload}"
    print(f"   rejected: {payload['message']}")


def step_empty_month():
    response = export(AUTH["faculty_a"], month=EMPTY_MONTH)  # default format: xlsx
    assert_file(response, EMPTY_MONTH, "xlsx")
    sheets = xlsx_sheets(response)
    assert sheets.get("Attendance") == [HEADERS], f"unexpected Attendance sheet: {sheets.get('Attendance')}"
    assert sheets.get("Summary") == [SUMMARY_HEADERS], f"unexpected Summary sheet: {sheets.get('Summary')}"
    response = export(AUTH["faculty_a"], month=EMPTY_MONTH, format="csv")
    assert_file(response, EMPTY_MONTH, "csv")
    assert csv_rows(response) == [HEADERS], f"unexpected csv: {csv_rows(response)}"
    print("   xlsx and csv hold the header rows only")


def step_timetable_classes_hod():
    """GET /timetable/classes: every occurrence, newest start time first."""
    status, payload = api("GET", "/timetable/classes", token=AUTH["hod"])
    assert status == 200, f"/timetable/classes returned {status}: {payload}"
    entries = payload["classes"]
    assert payload["count"] == len(entries), f"count does not match the list: {payload['count']}"
    assert len(entries) <= 500, f"listing is not capped at 500 rows: {len(entries)}"
    starts = [datetime.fromisoformat(entry["startTime"]) for entry in entries]
    assert starts == sorted(starts, reverse=True), "classes are not newest start time first"
    ids = {entry["id"] for entry in entries}
    if len(entries) < 500:
        # Not capped, so every class in the database must be listed. At the
        # cap, this run's (deliberately dated) rows may fall outside the
        # newest 500 - the filtered step below still proves they are there.
        missing = [class_id for class_id in STATE["class_ids"] if class_id not in ids]
        assert not missing, f"classes missing from the listing: {missing}"
    print(f"   {len(entries)} classes listed, newest start time first")


def step_timetable_classes_filter():
    """?professor_id= limits the listing, and each row names its professor."""
    status, payload = api(
        "GET", "/timetable/classes", token=AUTH["hod"],
        params={"professor_id": STATE["faculty_a_id"]},
    )
    assert status == 200, f"filtered listing returned {status}: {payload}"
    by_id = {entry["id"]: entry for entry in payload["classes"]}
    assert set(by_id) == {STATE["a1_id"], STATE["class_ids"][1]}, (
        f"faculty A's filtered classes are {sorted(by_id)}, expected {STATE['class_ids'][:2]}"
    )
    a1 = by_id[STATE["a1_id"]]
    assert set(a1) == CLASS_FIELDS, f"unexpected fields: {sorted(a1)}"
    assert a1["courseCode"] == COURSE_CODE, f"unexpected course code: {a1}"
    assert a1["room"] == "EXPORT-A1", f"unexpected room: {a1}"
    assert a1["professorId"] == STATE["faculty_a_id"], f"unexpected professor: {a1}"
    assert a1["professorName"] == FACULTY_A_NAME, f"unexpected professor name: {a1}"

    status, payload = api(
        "GET", "/timetable/classes", token=AUTH["hod"],
        params={"professor_id": STATE["faculty_b_id"]},
    )
    assert status == 200, f"filtered listing returned {status}: {payload}"
    entries = payload["classes"]
    ids = [entry["id"] for entry in entries]
    assert STATE["b1_id"] in ids, f"faculty B's class {STATE['b1_id']} is missing: {ids}"
    assert STATE["a1_id"] not in ids, f"faculty A's class leaked into the filtered list: {ids}"
    others = [entry for entry in entries if entry["professorId"] != STATE["faculty_b_id"]]
    assert not others, f"filter returned other professors' classes: {others}"
    assert entries[0]["professorName"] == FACULTY_B_NAME, f"unexpected professor name: {entries[0]}"
    print(f"   faculty A: {len(by_id)} classes; faculty B: {len(entries)} classes")


def step_timetable_classes_guards():
    for who, token in (("faculty", AUTH["faculty_a"]), ("student", AUTH["student_1"])):
        status, payload = api("GET", "/timetable/classes", token=token)
        assert status == 403, f"{who}: expected 403, got {status}: {payload}"
        assert payload.get("check") == "auth", f"{who}: unexpected check: {payload}"
        assert payload.get("error") == "insufficient_role", f"{who}: unexpected error: {payload}"
    status, payload = api("GET", "/timetable/classes")
    assert status == 401, f"no token: expected 401, got {status}: {payload}"
    assert payload.get("check") == "auth", f"no token: unexpected check: {payload}"
    print(f"   faculty and student rejected; no token -> 401 {payload.get('error')}")


def step_export_professor_scope():
    """HOD ?professor_id= exports that faculty member's classes only."""
    response = export(
        AUTH["hod"], month=TEST_MONTH, format="csv", professor_id=STATE["faculty_a_id"]
    )
    assert_file(response, TEST_MONTH, "csv")
    rows = csv_rows(response)
    assert rows[0] == HEADERS, f"unexpected headers: {rows[0]}"
    assert rows[1:] == STATE["expected"], (
        f"unexpected faculty A rows:\n   got      {rows[1:]}\n   expected {STATE['expected']}"
    )
    assert not [row for row in rows[1:] if row[2] == "EXPORT-B1"], (
        f"faculty B's class leaked into A's export: {rows[1:]}"
    )

    response = export(
        AUTH["hod"], month=TEST_MONTH, format="csv", professor_id=STATE["faculty_b_id"]
    )
    assert_file(response, TEST_MONTH, "csv")
    rows = csv_rows(response)
    assert rows[1:] == STATE["expected_b1"], (
        f"unexpected faculty B rows:\n   got      {rows[1:]}\n   expected {STATE['expected_b1']}"
    )
    print(f"   A: {len(STATE['expected'])} rows; B: {len(rows) - 1} row(s), no overlap")


def step_export_professor_conflicts():
    """A class of another professor, and bad ids, are rejected."""
    response = export(
        AUTH["hod"], month=TEST_MONTH, class_id=STATE["a1_id"], professor_id=STATE["faculty_b_id"]
    )
    payload = error_of(response)
    assert response.status_code == 400, f"expected 400, got {response.status_code}: {payload}"
    assert payload.get("check") == "request", f"unexpected check: {payload}"
    assert payload.get("error") == "class_professor_mismatch", f"unexpected error: {payload}"
    mismatch = payload["message"]

    for professor_id in ("abc", "-1", "1.5", STATE["missing_user_id"]):
        response = export(AUTH["hod"], month=TEST_MONTH, professor_id=professor_id)
        payload = error_of(response)
        assert response.status_code == 400, (
            f"professor_id={professor_id!r}: expected 400, got {response.status_code}: {payload}"
        )
        assert payload.get("check") == "request", (
            f"professor_id={professor_id!r}: unexpected check: {payload}"
        )
        assert payload.get("error") == "invalid_professor_id", (
            f"professor_id={professor_id!r}: unexpected error: {payload}"
        )
    print(f"   rejected: {mismatch}; and 4 bad professor_id values")


def step_export_faculty_scope():
    """Faculty may name only their own id; it changes nothing for them."""
    response = export(
        AUTH["faculty_a"], month=TEST_MONTH, professor_id=STATE["faculty_b_id"]
    )
    payload = error_of(response)
    assert response.status_code == 403, f"expected 403, got {response.status_code}: {payload}"
    assert payload.get("check") == "professor", f"unexpected check: {payload}"
    assert payload.get("error") == "not_class_professor", f"unexpected error: {payload}"

    response = export(
        AUTH["faculty_a"], month=TEST_MONTH, format="csv", professor_id=STATE["faculty_a_id"]
    )
    assert_file(response, TEST_MONTH, "csv")
    rows = csv_rows(response)
    assert rows[1:] == STATE["expected"], (
        f"their own id changed the export:\n   got      {rows[1:]}\n   expected {STATE['expected']}"
    )
    print(f"   rejected: {payload['message']}; own id -> {len(rows) - 1} rows, unchanged")


# --------------------------------------------------------------------------
# Cleanup - only this run's rows
# --------------------------------------------------------------------------
def cleanup() -> str:
    removed = 0
    with APP.app_context():
        user_ids = [u.id for u in User.query.filter(User.email.in_(ALL_EMAILS)).all()]
        class_ids = [
            c.id for c in ScheduledClass.query.filter(ScheduledClass.course_code.in_(COURSE_CODES)).all()
        ]
        # Both workload sessions of A2 share the same class key, so the
        # month-boundary session is collected (and deleted) with the rest.
        class_keys = [str(class_id) for class_id in class_ids] + [f"workload-{class_id}" for class_id in class_ids]
        session_ids = [
            s.id for s in AttendanceSession.query.filter(AttendanceSession.class_id.in_(class_keys)).all()
        ]
        removed += AttendanceCheckIn.query.filter(AttendanceCheckIn.session_id.in_(session_ids)).delete(
            synchronize_session=False,
        )
        removed += AttendanceSession.query.filter(AttendanceSession.id.in_(session_ids)).delete(
            synchronize_session=False,
        )
        removed += ScheduledClass.query.filter(ScheduledClass.id.in_(class_ids)).delete(
            synchronize_session=False,
        )
        for model in (FacultyProfile, StudentProfile):
            removed += model.query.filter(model.user_id.in_(user_ids)).delete(synchronize_session=False)
        removed += User.query.filter(User.id.in_(user_ids)).delete(synchronize_session=False)
        removed += InstitutionWhitelist.query.filter_by(email=HOD_EMAIL).delete(
            synchronize_session=False,
        )
        db.session.commit()
        assert User.query.filter(User.email.in_(ALL_EMAILS)).count() == 0
        assert ScheduledClass.query.filter(ScheduledClass.course_code.in_(COURSE_CODES)).count() == 0
        assert InstitutionWhitelist.query.filter_by(email=HOD_EMAIL).count() == 0
    return f"removed {removed} local rows"


# --------------------------------------------------------------------------
# Runner
# --------------------------------------------------------------------------
STEPS = [
    ("Preflight: backend /health", step_preflight),
    ("Auth setup: unique faculty A + B, an HOD and two students", step_auth_setup),
    ("Seed classes, sessions (attendance + workload) and check-ins", step_seed),
    ("Faculty A xlsx, all classes: headers + expected rows", step_xlsx_all_classes),
    ("Summary sheet: lectures per student and courses covered", step_summary_sheet),
    ("csv holds the same rows as the xlsx", step_csv_same_rows),
    ("class_id limits the export to one class", step_single_class),
    ("Month is the local calendar month", step_local_month_boundary),
    ("Faculty B exports faculty A's class -> 403", step_other_faculty_forbidden),
    ("Student exports -> 403", step_student_forbidden),
    ("Bad month -> 400 invalid_month", step_bad_month),
    ("Bad format -> 400 invalid_format", step_bad_format),
    ("Empty month -> headers only", step_empty_month),
    ("HOD lists every scheduled class", step_timetable_classes_hod),
    ("professor_id filters the class listing", step_timetable_classes_filter),
    ("Class listing: faculty/student -> 403, no token -> 401", step_timetable_classes_guards),
    ("HOD professor_id scopes the export to one faculty member", step_export_professor_scope),
    ("Mismatched class / bad professor_id -> 400", step_export_professor_conflicts),
    ("Faculty professor_id: other -> 403, own -> unchanged", step_export_faculty_scope),
]


def main() -> int:
    print("=" * 60)
    print("   MONTHLY ATTENDANCE EXPORT TESTS (test_export.py)")
    print("=" * 60)

    failed = None
    for index, (title, action) in enumerate(STEPS):
        print(f"\n{index}. {title}")
        try:
            action()
        except requests.RequestException as exc:
            print(f"   \u274c FAIL: cannot reach the backend at {BASE_URL} ({exc})")
            failed = (index, title)
            break
        except AssertionError as exc:
            print(f"   \u274c FAIL: {exc}")
            failed = (index, title)
            break
        except Exception as exc:
            print(f"   \u274c FAIL: unexpected {type(exc).__name__}: {exc}")
            failed = (index, title)
            break
        print("   \u2705 PASS")

    print("\nCleanup")
    cleanup_ok = True
    try:
        print(f"   \u2705 PASS: {cleanup()}")
    except Exception as exc:
        cleanup_ok = False
        print(f"   \u274c FAIL: cleanup failed: {type(exc).__name__}: {exc}")

    print("\n" + "=" * 60)
    if failed is not None:
        print(f"   TEST RUN FAILED AT STEP {failed[0]}: {failed[1]}")
    elif not cleanup_ok:
        print("   TEST RUN FAILED DURING CLEANUP")
    else:
        print(f"   ALL EXPORT TESTS PASSED ({len(STEPS)} steps)")
    print("=" * 60)
    return 0 if failed is None and cleanup_ok else 1


if __name__ == "__main__":
    sys.exit(main())
